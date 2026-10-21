#!/usr/bin/env python3
"""
╔══════════════════════════════════════════════════════╗
║         🍺  SERVEUR BAR LOCAL  v4.1                  ║
║         ✓ Stockage multi-fichiers JSON               ║
║         ✓ localStorage + localhost supportés         ║
║         ✓ Migration automatique depuis bar_data.json ║
║         ✓ Export Excel auto après modif membres      ║
║         ✓ Détection venv / Python système            ║
╚══════════════════════════════════════════════════════╝
"""

import json, os, socket, subprocess, sys, threading, time as _time
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
os.chdir(SCRIPT_DIR)

# ─── Redémarrage du serveur d'impression (port 8088) ──────────────────────────
IMPRESSION_PORT = 8088

def _kill_port(port):
    """Tue tout process qui écoute sur le port donné."""
    try:
        if sys.platform == "win32":
            result = subprocess.run(f'netstat -aon | findstr :{port} | findstr LISTENING',
                                    shell=True, capture_output=True, text=True)
            for line in result.stdout.strip().splitlines():
                parts = line.split()
                if parts:
                    subprocess.run(['taskkill', '/F', '/PID', parts[-1]], capture_output=True)
        else:
            subprocess.run(["fuser", "-k", f"{port}/tcp"], capture_output=True)
    except Exception:
        pass

def restart_serveur_impression():
    """Tue puis relance serveur_impression.py. Retourne (ok, message)."""
    impression_py = None
    for name in ["serveur_impression.py", "serveur_impression_v3TVA.py"]:
        candidate = os.path.join(SCRIPT_DIR, name)
        if os.path.exists(candidate):
            impression_py = candidate
            break
    if not impression_py:
        return False, "serveur_impression.py introuvable"

    _kill_port(IMPRESSION_PORT)
    _time.sleep(0.5)

    try:
        py = sys.executable or "python"
        kwargs = {"cwd": SCRIPT_DIR}
        if sys.platform == "win32":
            kwargs["creationflags"] = subprocess.CREATE_NEW_CONSOLE
        p = subprocess.Popen([py, impression_py], **kwargs)
        return True, f"Relancé (PID {p.pid})"
    except Exception as e:
        return False, str(e)

# ─── Détection venv ───────────────────────────────────────────────────────────
IN_VENV = (
    hasattr(sys, "real_prefix")                          # virtualenv classique
    or (hasattr(sys, "base_prefix") and sys.base_prefix != sys.prefix)  # venv stdlib
)

print()
print("  🐍 Python      :", sys.executable)
print("  📦 Environnement:", "✅ venv actif" if IN_VENV else "⚠️  Python système (pas de venv)")
print("  📂 Dossier     :", SCRIPT_DIR)

# ─── Dossier sauvegarde Excel ─────────────────────────────────────────────────
SAUVEGARDE_DIR = os.path.join(SCRIPT_DIR, "sauvegarde données")
os.makedirs(SAUVEGARDE_DIR, exist_ok=True)

# ─── Import openpyxl avec message clair ──────────────────────────────────────
try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_OK = True
    print("  📊 openpyxl    : ✅ disponible (", openpyxl.__version__, ")")
except ImportError:
    EXCEL_OK = False
    print("  📊 openpyxl    : ❌ NON installé")
    if IN_VENV:
        print("     → Lance : pip install openpyxl  (dans ton venv)")
    else:
        print("     → Lance : pip install openpyxl")
    print("     → Les exports/sauvegardes Excel seront désactivés")

print()

# ─── Fichiers JSON séparés ────────────────────────────────────────────────────
FILES = {
    "membres":  os.path.join(SCRIPT_DIR, "membres.json"),
    "tickets":  os.path.join(SCRIPT_DIR, "tickets.json"),
    "articles": os.path.join(SCRIPT_DIR, "articles.json"),
    "tables":   os.path.join(SCRIPT_DIR, "tables.json"),
    "rapports": os.path.join(SCRIPT_DIR, "rapports.json"),
    "stock":    os.path.join(SCRIPT_DIR, "stock.json"),
    "settings": os.path.join(SCRIPT_DIR, "settings.json"),
    "misc":     os.path.join(SCRIPT_DIR, "bar_data.json"),
}

FILE_KEYS = {
    "membres":  ["membres", "lastMembreId", "memberPins", "membresLogoOverride"],
    "tickets":  ["tickets", "ticketsHistorique", "deletedTickets", "refundedTickets",
                 "availableTicketIds", "lastTicketId", "lastTicketNumero"],
    "articles": ["articlesData", "articles", "tvaSettings", "lastArticleId", "lastCategorieId"],
    "tables":   ["tables", "reservations"],
    "rapports": ["rapports", "corbeilleRapports", "numX", "numZ", "lastZCloture",
                 "ctrl_log", "ctrl_rapports"],
    "stock":    ["stock", "stockData", "stockHistory"],
    "settings": ["companySettings", "displayOptions", "users", "loginHistory",
                 "lastLoginISO", "serverConfig", "emailSubject", "emailBody", "adminPassword"],
    "misc":     ["pendingOrders", "asbl_gestion_v1", "mailSubject", "mailBody",
                 "historiqueEmails", "historiqueFactures"],
}

KEY_TO_FILE = {}
for _fname, _keys in FILE_KEYS.items():
    for _k in _keys:
        KEY_TO_FILE[_k] = _fname

ALL_KEYS = [k for keys in FILE_KEYS.values() for k in keys]

# ─── Locks par fichier ────────────────────────────────────────────────────────
FILE_LOCKS = {fname: threading.Lock() for fname in FILES}

DEFAULTS = {
    "membres":  {"membres": [], "lastMembreId": 0, "memberPins": {}, "membresLogoOverride": None},
    "tickets":  {"tickets": [], "ticketsHistorique": [], "deletedTickets": [],
                 "refundedTickets": [], "availableTicketIds": [], "lastTicketId": 0,
                 "lastTicketNumero": 0},
    "articles": {"articlesData": {"categories": [], "articles": []}, "articles": [],
                 "tvaSettings": {}, "lastArticleId": 0, "lastCategorieId": 0},
    "tables":   {"tables": [], "reservations": []},
    "rapports": {"rapports": [], "corbeilleRapports": [], "numX": 0, "numZ": 0,
                 "lastZCloture": None, "ctrl_log": [], "ctrl_rapports": []},
    "stock":    {"stock": {}, "stockData": {}, "stockHistory": []},
    "settings": {"companySettings": {}, "displayOptions": {}, "users": [],
                 "loginHistory": [], "lastLoginISO": None, "serverConfig": {},
                 "emailSubject": "Votre carte de membre", "emailBody": "", "adminPassword": ""},
    "misc":     {"pendingOrders": [], "asbl_gestion_v1": {"dons": [], "inventaire": [], "settings": {}},
                 "mailSubject": "", "mailBody": "", "historiqueEmails": [], "historiqueFactures": []},
}

# ─── Lecture / écriture JSON ──────────────────────────────────────────────────
def load_file(fname):
    path = FILES[fname]
    defaults = dict(DEFAULTS[fname])
    if not os.path.exists(path):
        legacy = os.path.join(SCRIPT_DIR, "bar_data.json")
        if os.path.exists(legacy) and fname != "misc":
            try:
                with open(legacy, "r", encoding="utf-8") as f:
                    ld = json.load(f)
                migrated = {k: ld[k] for k in FILE_KEYS[fname] if k in ld}
                if migrated:
                    data = {**defaults, **migrated}
                    save_file(fname, data)
                    print(f"  📦 Migration → {os.path.basename(path)} ({len(migrated)} clés)")
                    return data
            except Exception:
                pass
        return defaults
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for k, v in defaults.items():
            # Corrige aussi le cas où la clé existe mais vaut null (JSON) → None (Python) :
            # ça arrivait par ex. avec "rapports": null suite à une sync corrompue,
            # et ça faisait planter le serveur plus loin (len(None) → crash → sync bloquée).
            if k not in data or data[k] is None:
                data[k] = v
        return data
    except Exception as e:
        print(f"  ⚠️  Erreur lecture {os.path.basename(path)}: {e}")
        return defaults

def save_file(fname, data):
    path = FILES[fname]
    clean = {k: data[k] for k in FILE_KEYS[fname] if k in data}
    with FILE_LOCKS[fname]:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(clean, f, ensure_ascii=False, indent=2)

def load_all():
    merged = {}
    for fname in FILES:
        merged.update(load_file(fname))
    return merged

def save_keys(data, keys_changed):
    files_to_save = set()
    for k in keys_changed:
        files_to_save.add(KEY_TO_FILE.get(k, "misc"))
    for fname in files_to_save:
        save_file(fname, data)

def reset_all_data():
    """Remet TOUTES les données à zéro : les fichiers JSON ET les sauvegardes Excel
    auto-générées dans "sauvegarde données/" (sauvegarde_courante_*.xlsx), qui sinon
    resteraient affichées avec les anciennes données jusqu'au prochain export."""
    for fname in FILES:
        save_file(fname, dict(DEFAULTS[fname]))
    for xlsx_name in ("sauvegarde_courante_membres.xlsx",
                       "sauvegarde_courante_tickets.xlsx",
                       "sauvegarde_courante_rapports.xlsx"):
        xlsx_path = os.path.join(SAUVEGARDE_DIR, xlsx_name)
        try:
            if os.path.exists(xlsx_path):
                os.remove(xlsx_path)
        except Exception as e:
            print(f"  ⚠️  Impossible de supprimer {xlsx_name} : {e}")

# ─── Fusion anti-résurrection (tickets/tables) ────────────────────────────────
# Problème résolu ici : avec plusieurs postes (tablette + PC), l'ancien /api/sync
# remplaçait tout le tableau "tickets" par celui envoyé par le poste qui synchronisait
# en dernier. Si un poste avait une copie en mémoire pas encore rafraîchie (donc encore
# avec un ticket déjà payé/fermé sur un autre poste), ce ticket "fermé" repartait dans
# le tableau des tickets ouverts avec ses anciennes boissons → c'est le ticket qui
# "réapparaît" sur une table qu'on vient de repointer.
# Fix : on fusionne par id (union, jamais de perte), et surtout on ne laisse jamais un
# ticket réapparaître comme "ouvert" s'il existe déjà dans l'historique/la corbeille/les
# remboursements (= c'est un ticket clos, il ne doit plus jamais être un ticket vivant).

def _merge_list_by_id(disk_list, body_list):
    """Fusionne 2 listes de dicts par id : union des ids, la version du body gagne
    en cas de conflit (c'est la plus fraîche puisqu'envoyée par le poste actif)."""
    merged = {}
    for item in (disk_list or []):
        if isinstance(item, dict) and "id" in item:
            merged[str(item["id"])] = item
    for item in (body_list or []):
        if isinstance(item, dict) and "id" in item:
            merged[str(item["id"])] = item
    return list(merged.values())

def _reconcile_ticket_buckets(data):
    """Un ticket payé ne doit jamais exister dans 2 des 3 listes en même temps
    (ticketsHistorique / deletedTickets / refundedTickets). Sans ça : si tu fais un
    remboursement sur un poste, et qu'un autre poste pas encore rafraîchi resynchronise
    sa vieille copie de ticketsHistorique (qui contient encore ce ticket), il redevient
    "actif" dans ticketsHistorique en plus d'être dans refundedTickets → il serait recompté
    en double dans le prochain recalcul de rapport Z.
    Priorité en cas de conflit : refundedTickets > deletedTickets > ticketsHistorique
    (un remboursement ou une suppression explicite ne doit jamais être écrasé par une
    copie périmée d'un autre poste qui n'était pas encore au courant)."""
    priority = ["refundedTickets", "deletedTickets", "ticketsHistorique"]
    by_bucket = {k: {} for k in priority}
    for k in priority:
        for item in data.get(k, []) or []:
            if isinstance(item, dict) and "id" in item:
                by_bucket[k][str(item["id"])] = item

    winner_bucket, winner_item = {}, {}
    for k in priority:
        for id_, item in by_bucket[k].items():
            if id_ not in winner_bucket:
                winner_bucket[id_] = k
                winner_item[id_] = item

    for k in priority:
        data[k] = [winner_item[id_] for id_ in winner_bucket if winner_bucket[id_] == k]

    return set(winner_bucket.keys())

def merge_tickets_and_tables(data, body):
    """Fusionne tickets/tables/historique/corbeille/remboursements de façon sûre.
    Retourne la liste des clés effectivement modifiées."""
    changed = []

    hist_keys = ["ticketsHistorique", "deletedTickets", "refundedTickets"]
    for k in hist_keys:
        if k in body:
            data[k] = _merge_list_by_id(data.get(k), body[k])
            changed.append(k)

    # Chaque ticket ne doit exister que dans UNE seule des 3 listes (voir docstring)
    dead_ids = _reconcile_ticket_buckets(data)

    if "tickets" in body:
        merged = _merge_list_by_id(data.get("tickets"), body["tickets"])
        merged = [t for t in merged if str(t.get("id")) not in dead_ids]
        data["tickets"] = merged
        changed.append("tickets")

    if "tables" in body:
        merged_tables = _merge_list_by_id(data.get("tables"), body["tables"])
        # Une table qui pointe encore vers un ticket clos est libérée automatiquement
        for t in merged_tables:
            tid = t.get("ticketId")
            if tid is not None and str(tid) in dead_ids:
                t["ticketId"] = None
        data["tables"] = merged_tables
        changed.append("tables")

    return changed

client_display_data = None

# ─── Helpers Excel ────────────────────────────────────────────────────────────
def style_header(ws, row=1):
    if not EXCEL_OK:
        return
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")

def make_wb_rapports(rapports):
    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé Rapports"
    headers = ["N° Rapport","Type","Date Clôture","Heure Clôture","Date Ouverture",
               "Heure Ouverture","Utilisateur","CA Net (€)","CA Global (€)",
               "Total Espèces (€)","Total Banque (€)","Fond de Caisse (€)",
               "Total Global (€)","Nb Tickets","Nb Tickets Actifs","Articles Vendus",
               "Valeur Offerts (€)","Total Dons (€)","Dons Espèces (€)","Dons Banque (€)"]
    ws.append(headers)
    for r in rapports:
        ws.append([
            f"{r.get('type','Z')}-{r.get('numero','')}",
            r.get('type','Z'), r.get('dateCloture',''), r.get('heureCloture',''),
            r.get('dateOuverture',''), r.get('heureOuverture',''), r.get('utilisateur',''),
            float(r.get('ca',0)), float(r.get('caGlobal', r.get('ca',0))),
            float(r.get('totalEspeces',0)), float(r.get('totalBanque',0)),
            float(r.get('fondCaisse',0)), float(r.get('totalGlobal',0)),
            int(r.get('nbTickets',0)), int(r.get('nbTicketsActifs',0)),
            int(r.get('totalArticles',0)), float(r.get('totalOfferts',0)),
            float(r.get('totalDons',0)), float(r.get('donsEspeces',0)),
            float(r.get('donsBanque',0))
        ])
    style_header(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    wt = wb.create_sheet("Ventilation TVA")
    wt.append(["N° Rapport","Date","Taux TVA (%)","Base HTVA (€)","Montant TVA (€)","Total TVAC (€)"])
    for r in rapports:
        tva = r.get('tvaVentilation', {})
        for taux, v in sorted(tva.items(), key=lambda x: float(x[0])):
            wt.append([f"{r.get('type','Z')}-{r.get('numero','')}", r.get('dateCloture',''),
                       float(taux), float(v.get('base',0)), float(v.get('tva',0)), float(v.get('total',0))])
    style_header(wt)

    wa = wb.create_sheet("Articles par Catégorie")
    wa.append(["N° Rapport","Date","Catégorie","Article","Quantité","Total (€)"])
    for r in rapports:
        cats = r.get('categories', {})
        for cat, arts in cats.items():
            for art, obj in arts.items():
                wa.append([f"{r.get('type','Z')}-{r.get('numero','')}", r.get('dateCloture',''),
                           cat, art, int(obj.get('qte',0)), float(obj.get('total',0))])
    style_header(wa)
    return wb

def make_wb_tickets(tickets):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append(["ID","N° Ticket","Date","Heure","Mode Paiement","Total (€)",
               "Espèces (€)","Banque (€)","Table","Statut","Remboursé","Supprimé","Articles"])
    for t in tickets:
        arts = t.get('articles', [])
        arts_str = " | ".join(
            f"{a.get('nom', a.get('name',''))} x{a.get('quantite', a.get('qty',1))} ({float(a.get('prixTotal', a.get('prix',0))):.2f}€)"
            for a in arts
        ) if arts else ""
        ws.append([
            str(t.get('id','')), str(t.get('numero', t.get('id',''))),
            t.get('datePaiement', t.get('date','')), t.get('heurePaiement', t.get('heure','')),
            t.get('modePaiement',''), float(t.get('total',0)),
            float(t.get('montantEspeces',0)), float(t.get('montantBanque',0)),
            t.get('table', t.get('tableNom','')), t.get('statut',''),
            'Oui' if t.get('refundedAt') else 'Non',
            'Oui' if t.get('deletedAt') else 'Non',
            arts_str
        ])
    style_header(ws)
    ws.column_dimensions['M'].width = 40

    wd = wb.create_sheet("Détail Articles")
    wd.append(["ID Ticket","N° Ticket","Date","Article","Catégorie","Quantité","Prix unitaire (€)","Total ligne (€)","Offert"])
    for t in tickets:
        for a in t.get('articles', []):
            wd.append([
                str(t.get('id','')), str(t.get('numero', t.get('id',''))),
                t.get('datePaiement', t.get('date','')),
                a.get('nom', a.get('name','')), a.get('categorie', a.get('category','')),
                int(a.get('quantite', a.get('qty',1))),
                float(a.get('prix', a.get('prixUnitaire',0))),
                float(a.get('prixTotal', a.get('prix',0) * a.get('quantite',1))),
                'Oui' if a.get('offert') else 'Non'
            ])
    style_header(wd)
    return wb

def make_wb_membres(membres):
    wb = Workbook()
    ws = wb.active
    ws.title = "Membres"
    ws.append(["ID","Prénom","Nom","Adresse","Localité","Email","Téléphone","Statut","Date Adhésion"])
    for m in membres:
        ws.append([str(m.get('id','')), m.get('prenom',''), m.get('nom',''),
                   m.get('adresse',''), m.get('localite',''), m.get('email',''),
                   m.get('telephone',''), m.get('statut',''), m.get('dateAdhesion','')])
    style_header(ws)
    for col, w in zip('ABCDEFGHI', [8,14,14,22,16,26,14,14,14]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Statistiques")
    actifs = sum(1 for m in membres if str(m.get('statut','')).lower() == 'actif')
    ws2.append(["Statistique","Valeur"])
    ws2.append(["Total membres", len(membres)])
    ws2.append(["Membres actifs", actifs])
    ws2.append(["Membres inactifs", len(membres)-actifs])
    import datetime
    ws2.append(["Export le", datetime.date.today().strftime('%d/%m/%Y')])
    style_header(ws2)
    return wb

def wb_to_bytes(wb):
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ─── Export Excel automatique (silencieux, appelé après chaque modif membre) ──
def auto_export_membres():
    if not EXCEL_OK:
        print("  ⚠️  Export Excel membres ignoré (openpyxl non disponible)")
        return
    try:
        membres = load_file("membres").get("membres", [])
        wb = make_wb_membres(membres)
        save_path = os.path.join(SAUVEGARDE_DIR, "sauvegarde_courante_membres.xlsx")
        wb.save(save_path)
        print(f"  💾 Sauvegarde Excel membres → {len(membres)} membres enregistrés")
    except Exception as e:
        print(f"  ⚠️  Erreur sauvegarde Excel membres : {e}")

# ─── Présences actives (en mémoire) ──────────────────────────────────────────
_presence = {}
_presence_lock = threading.Lock()
PRESENCE_TTL = 30

def presence_get():
    now = _time.time()
    with _presence_lock:
        for page in list(_presence.keys()):
            _presence[page] = {d: i for d, i in _presence[page].items()
                               if now - i["ts"] < PRESENCE_TTL}
            if not _presence[page]:
                del _presence[page]
        return {page: list(v.values()) for page, v in _presence.items()}

def presence_set(page, device_id, user):
    with _presence_lock:
        if user:
            _presence.setdefault(page, {})[device_id] = {
                "user": user, "page": page, "ts": _time.time()
            }
        else:
            if page in _presence:
                _presence[page].pop(device_id, None)
                if not _presence[page]:
                    del _presence[page]

def data_is_empty():
    d = load_all()
    return (len(d.get("articlesData", {}).get("articles", [])) == 0
            and len(d.get("tables", [])) == 0)

# ─── Handler HTTP ─────────────────────────────────────────────────────────────
class BarHandler(SimpleHTTPRequestHandler):

    def log_message(self, fmt, *args):
        status = args[1] if len(args) > 1 else "?"
        icon = "✅" if str(status).startswith("2") else "⚠️ "
        print(f"  {icon} {self.command:6} {self.path.split('?')[0]}")

    def get_cors_origin(self):
        origin = self.headers.get("Origin", "")
        return origin if origin else "*"

    def send_json(self, code, obj):
        body = json.dumps(obj, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", self.get_cors_origin())
        self.send_header("Access-Control-Allow-Credentials", "true")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()
        self.wfile.write(body)

    def read_body(self):
        n = int(self.headers.get("Content-Length", 0))
        if n == 0:
            return {}
        try:
            return json.loads(self.rfile.read(n))
        except:
            return {}

    def do_OPTIONS(self):
        origin = self.headers.get("Origin", "")
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", origin if origin else "*")
        self.send_header("Access-Control-Allow-Credentials", "true")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    # ── GET ───────────────────────────────────────────────────────────────────
    def do_GET(self):
        path = urlparse(self.path).path

        if path in ("/", ""):
            self.send_response(302)
            self.send_header("Location", "/tables.html")
            self.end_headers()
            return

        if path == "/api/data":
            self.send_json(200, load_all())
            return

        if path == "/api/membres":
            self.send_json(200, load_file("membres").get("membres", []))
            return

        if path == "/api/tables":
            self.send_json(200, load_file("tables").get("tables", []))
            return

        if path == "/api/articles":
            self.send_json(200, load_file("articles").get("articlesData", {"categories":[], "articles":[]}))
            return

        if path == "/api/tickets":
            d = load_file("tickets")
            self.send_json(200, {"tickets": d.get("tickets", []),
                                 "ticketsHistorique": d.get("ticketsHistorique", [])})
            return

        if path == "/api/rapports":
            self.send_json(200, load_file("rapports"))
            return

        if path == "/api/stock":
            self.send_json(200, load_file("stock"))
            return

        if path == "/api/settings":
            self.send_json(200, load_file("settings"))
            return

        if path == "/api/client-display":
            global client_display_data
            if client_display_data is None:
                self.send_json(200, {"empty": True})
            else:
                self.send_json(200, client_display_data)
            return

        if path == "/api/pending":
            self.send_json(200, load_file("misc").get("pendingOrders", []))
            return

        if path == "/api/presence":
            self.send_json(200, presence_get())
            return

        if path == "/api/status":
            self.send_json(200, {
                "ok": True,
                "excel_ok": EXCEL_OK,
                "in_venv": IN_VENV,
                "python": sys.executable,
                "empty": data_is_empty(),
                "tables":   len(load_file("tables").get("tables", [])),
                "membres":  len(load_file("membres").get("membres", [])),
                "articles": len(load_file("articles").get("articlesData", {}).get("articles", [])),
                "rapports": len(load_file("rapports").get("rapports", [])),
                "pending":  len(load_file("misc").get("pendingOrders", [])),
            })
            return

        if path in ("/api/import-excel-rapports", "/api/import-excel-tickets", "/api/import-excel-membres"):
            self.handle_excel_import_post(path)
            return

        if path.startswith("/api/"):
            self.send_json(404, {"error": f"Route inconnue: {path}"})
            return

        super().do_GET()

    # ── POST ──────────────────────────────────────────────────────────────────
    def do_POST(self):
        try:
            self._do_POST_inner()
        except Exception as e:
            import traceback
            print("  🔥 ERREUR do_POST (voir détail ci-dessous) :")
            traceback.print_exc()
            try:
                self.send_json(500, {"ok": False, "error": str(e)})
            except Exception:
                # La connexion est peut-être déjà cassée, on ne peut plus répondre
                pass

    def _do_POST_inner(self):
        path = urlparse(self.path).path
        body = self.read_body()

        if path == "/api/client-display":
            global client_display_data
            client_display_data = body
            self.send_json(200, {"ok": True})
            return

        if path == "/api/presence":
            presence_set(
                body.get("page", "?"),
                body.get("deviceId", "anon"),
                body.get("user")
            )
            self.send_json(200, {"ok": True, "presence": presence_get()})
            return

        if path == "/api/sync":
            data = load_all()
            changed_keys = []

            # tickets/tables/historique/corbeille/remboursements : fusion sécurisée
            # (voir merge_tickets_and_tables) au lieu d'un simple écrasement, pour ne
            # jamais laisser un ticket déjà fermé "ressusciter" avec ses anciennes boissons.
            special_keys = {"tickets", "tables", "ticketsHistorique", "deletedTickets", "refundedTickets"}
            try:
                changed_keys.extend(merge_tickets_and_tables(data, body))
            except Exception as e:
                import traceback
                print("  🔥 ERREUR dans merge_tickets_and_tables :")
                traceback.print_exc()
                self.send_json(500, {"ok": False, "error": f"merge_tickets_and_tables: {e}"})
                return

            for key in ALL_KEYS:
                if key in special_keys or key not in body:
                    continue
                val = body[key]
                disk_val = data.get(key)
                disk_empty = (disk_val in (None, [], {}, 0, ""))
                body_empty = (val in (None, [], {}, 0, ""))
                if not body_empty or disk_empty:
                    data[key] = val
                    changed_keys.append(key)
            if changed_keys:
                save_keys(data, changed_keys)
                nb_art = len((data.get("articlesData") or {}).get("articles") or [])
                nb_tab = len(data.get("tables") or [])
                nb_mem = len(data.get("membres") or [])
                nb_rap = len(data.get("rapports") or [])
                print(f"  💾 Sync → {nb_tab} tables · {nb_art} articles · {nb_mem} membres · {nb_rap} rapports")
                # Si des membres ont changé via sync, on exporte aussi
                if "membres" in changed_keys:
                    auto_export_membres()
            self.send_json(200, {"ok": True, "changed": len(changed_keys) > 0, **data})
            return

        # Sauvegardes ciblées par domaine
        for domain in ["tickets", "rapports", "stock", "settings", "articles", "tables"]:
            if path == f"/api/save/{domain}":
                d = load_file(domain)
                for k in FILE_KEYS[domain]:
                    if k in body:
                        d[k] = body[k]
                save_file(domain, d)
                self.send_json(200, {"ok": True})
                return

        if path == "/api/order":
            order = body
            if not order.get("id"):
                self.send_json(400, {"error": "Commande invalide"})
                return
            d = load_file("misc")
            d.setdefault("pendingOrders", []).append(order)
            save_file("misc", d)
            print(f"  🔔 Commande → Table {order.get('tableNom','?')} · {len(order.get('articles',[]))} article(s)")
            self.send_json(200, {"ok": True})
            return

        if path == "/api/order/accept":
            oid   = body.get("id")
            order = body.get("order", {})
            misc  = load_file("misc")
            misc["pendingOrders"] = [o for o in misc.get("pendingOrders", []) if o.get("id") != oid]
            if order and order.get("tableId"):
                tables_d  = load_file("tables")
                tickets_d = load_file("tickets")
                table = next((t for t in tables_d.get("tables", []) if t["id"] == order["tableId"]), None)
                if table:
                    ticket = next((t for t in tickets_d.get("tickets", []) if t["id"] == table.get("ticketId")), None)
                    if not ticket:
                        tickets_d["lastTicketId"]    = tickets_d.get("lastTicketId", 0) + 1
                        tickets_d["lastTicketNumero"] = tickets_d.get("lastTicketNumero", 0) + 1
                        ticket = {"id": tickets_d["lastTicketId"], "numero": tickets_d["lastTicketNumero"],
                                  "tableId": table["id"], "date": body.get("date",""),
                                  "heure": body.get("heure",""), "membre": None, "articles": []}
                        tickets_d.setdefault("tickets", []).append(ticket)
                        table["ticketId"] = ticket["id"]
                    for art in order.get("articles", []):
                        ex = next((a for a in ticket["articles"]
                                   if a.get("id")==art.get("id") and a.get("prix")==art.get("prix")
                                   and not a.get("offert")), None)
                        if ex:
                            ex["qte"] = ex.get("qte", 0) + art.get("qte", 1)
                        else:
                            ticket["articles"].append({"id": art.get("id"), "nom": art.get("nom"),
                                                       "prix": art.get("prix",0), "qte": art.get("qte",1),
                                                       "offert": False})
                    save_file("tickets", tickets_d)
                    save_file("tables", tables_d)
            save_file("misc", misc)
            print(f"  ✅ Commande acceptée · Table {order.get('tableNom','?')}")
            self.send_json(200, {"ok": True})
            return

        if path == "/api/restart-impression":
            ok, msg = restart_serveur_impression()
            print(f"  🖨️  Redémarrage serveur impression → {msg}")
            self.send_json(200 if ok else 500, {"ok": ok, "message": msg})
            return

        if path == "/api/order/refuse":
            oid  = body.get("id")
            misc = load_file("misc")
            misc["pendingOrders"] = [o for o in misc.get("pendingOrders", []) if o.get("id") != oid]
            save_file("misc", misc)
            print(f"  ❌ Commande refusée · ID {oid}")
            self.send_json(200, {"ok": True})
            return

        # ── Membre : ajout / modification ──────────────────────────────────
        if path == "/api/membre":
            membre = body
            mid = str(membre.get("id", ""))
            d = load_file("membres")
            membres = d.get("membres", [])
            idx = next((i for i, m in enumerate(membres) if str(m.get("id","")) == mid), None)
            if idx is not None:
                membres[idx] = membre
                print(f"  👤 Membre modifié : {membre.get('prenom','')} {membre.get('nom','')} (ID {mid})")
            else:
                d["lastMembreId"] = d.get("lastMembreId", 0) + 1
                membre["id"] = str(d["lastMembreId"])
                membres.append(membre)
                print(f"  👤 Nouveau membre : {membre.get('prenom','')} {membre.get('nom','')} (ID {membre['id']})")
            d["membres"] = membres
            save_file("membres", d)
            auto_export_membres()   # ✅ export Excel automatique
            self.send_json(200, {"ok": True, "membre": membre})
            return

        # ── Membre : suppression ───────────────────────────────────────────
        if path == "/api/membre/delete":
            mid = str(body.get("id", ""))
            d = load_file("membres")
            avant = len(d.get("membres", []))
            d["membres"] = [m for m in d.get("membres", []) if str(m.get("id","")) != mid]
            apres = len(d["membres"])
            save_file("membres", d)
            auto_export_membres()   # ✅ export Excel automatique
            print(f"  🗑️  Membre supprimé · ID {mid} ({avant} → {apres} membres)")
            self.send_json(200, {"ok": True})
            return

        if path == "/api/save":
            tables_d  = load_file("tables")
            tickets_d = load_file("tickets")
            stock_d   = load_file("stock")
            changed = []
            if "tables"  in body: tables_d["tables"]   = body["tables"];  changed.append("tables")
            if "tickets" in body: tickets_d["tickets"]  = body["tickets"]; changed.append("tickets")
            if "stock"   in body: stock_d["stock"]      = body["stock"];   changed.append("stock")
            if "tables"  in changed: save_file("tables", tables_d)
            if "tickets" in changed: save_file("tickets", tickets_d)
            if "stock"   in changed: save_file("stock", stock_d)
            self.send_json(200, {"ok": True})
            return

        # ── Remise à zéro complète (bouton "danger" dans paramètres.html) ──
        if path == "/api/reset-all":
            reset_all_data()
            print("  🧨 RESET TOTAL demandé depuis paramètres.html — toutes les données ont été remises à zéro")
            self.send_json(200, {"ok": True})
            return

        # ── Export Excel ───────────────────────────────────────────────────
        if path in ("/api/export-excel-rapports", "/api/export-excel-tickets", "/api/export-excel-membres"):
            if not EXCEL_OK:
                self.send_json(500, {"error": "openpyxl non installé. Lance : pip install openpyxl"})
                return
            qs = parse_qs(urlparse(self.path).query)
            is_auto = qs.get("auto", ["0"])[0] == "1"
            try:
                if path == "/api/export-excel-rapports":
                    rapports = body.get("rapports", [])
                    wb = make_wb_rapports(rapports)
                    fname_auto = "sauvegarde_courante_rapports.xlsx"
                    fname_dl   = "Rapports_Z.xlsx"
                elif path == "/api/export-excel-tickets":
                    tickets = body.get("tickets", [])
                    wb = make_wb_tickets(tickets)
                    fname_auto = "sauvegarde_courante_tickets.xlsx"
                    fname_dl   = "Historique_Tickets.xlsx"
                else:
                    membres = body.get("membres", [])
                    wb = make_wb_membres(membres)
                    fname_auto = "sauvegarde_courante_membres.xlsx"
                    fname_dl   = "Membres.xlsx"

                wb.save(os.path.join(SAUVEGARDE_DIR, fname_auto))
                if not is_auto:
                    print(f"  💾 Export manuel Excel → {fname_dl}")

                if is_auto:
                    self.send_json(200, {"ok": True})
                    return

                data_bytes = wb_to_bytes(wb)
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{fname_dl}"')
                self.send_header("Content-Length", str(len(data_bytes)))
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(data_bytes)
            except Exception as e:
                self.send_json(500, {"error": str(e)})
            return

        # ── Import Excel ───────────────────────────────────────────────────
        if path in ("/api/import-excel-rapports", "/api/import-excel-tickets", "/api/import-excel-membres"):
            self.handle_excel_import_post(path)
            return

        self.send_json(404, {"error": f"Route inconnue : {path}"})

    # ── Import Excel (reçoit bytes xlsx → parse → retour JSON) ───────────────
    def handle_excel_import_post(self, path):
        if not EXCEL_OK:
            self.send_json(500, {"error": "openpyxl non installé. Lance : pip install openpyxl"})
            return
        try:
            n = int(self.headers.get("Content-Length", 0))
            raw = self.rfile.read(n)
            import io
            wb = openpyxl.load_workbook(io.BytesIO(raw))

            if path in ("/api/import-excel-rapports", "/api/import-excel-rapports"):
                ws = wb["Résumé Rapports"]
                rows = list(ws.values)[1:]
                rapports = []
                for r in rows:
                    if not r[0]: continue
                    num = str(r[0]).split('-')[1] if '-' in str(r[0]) else str(r[0])
                    rap = {"numero": num, "type": r[1] or "Z",
                           "dateCloture": str(r[2] or ""), "heureCloture": str(r[3] or ""),
                           "dateOuverture": str(r[4] or ""), "heureOuverture": str(r[5] or ""),
                           "utilisateur": str(r[6] or ""),
                           "ca": float(r[7] or 0), "caGlobal": float(r[8] or 0),
                           "totalEspeces": float(r[9] or 0), "totalBanque": float(r[10] or 0),
                           "fondCaisse": float(r[11] or 0), "totalGlobal": float(r[12] or 0),
                           "nbTickets": int(r[13] or 0), "nbTicketsActifs": int(r[14] or 0),
                           "totalArticles": int(r[15] or 0), "totalOfferts": float(r[16] or 0),
                           "totalDons": float(r[17] or 0), "donsEspeces": float(r[18] or 0),
                           "donsBanque": float(r[19] or 0),
                           "categories": {}, "tvaVentilation": {}, "articlesOfferts": {}}
                    rapports.append(rap)
                idx = {f"{r['type']}-{r['numero']}": r for r in rapports}
                if "Ventilation TVA" in wb.sheetnames:
                    for row in list(wb["Ventilation TVA"].values)[1:]:
                        rap = idx.get(str(row[0]))
                        if rap and row[2]:
                            rap["tvaActif"] = True
                            rap["tvaVentilation"][str(row[2])] = {
                                "base": float(row[3] or 0), "tva": float(row[4] or 0), "total": float(row[5] or 0)}
                if "Articles par Catégorie" in wb.sheetnames:
                    for row in list(wb["Articles par Catégorie"].values)[1:]:
                        rap = idx.get(str(row[0]))
                        if rap and row[2] and row[3]:
                            cat, art = str(row[2]), str(row[3])
                            rap["categories"].setdefault(cat, {})[art] = {
                                "qte": int(row[4] or 0), "total": float(row[5] or 0)}
                self.send_json(200, {"success": True, "count": len(rapports), "rapports": rapports})

            elif path == "/api/import-excel-tickets":
                ws = wb["Tickets"]
                rows = list(ws.values)[1:]
                tickets = [{"id": str(r[0] or ""), "numero": str(r[1] or r[0] or ""),
                            "datePaiement": str(r[2] or ""), "heurePaiement": str(r[3] or ""),
                            "modePaiement": str(r[4] or ""), "total": float(r[5] or 0),
                            "montantEspeces": float(r[6] or 0), "montantBanque": float(r[7] or 0),
                            "table": str(r[8] or ""), "statut": str(r[9] or ""), "articles": []}
                           for r in rows if r[0]]
                self.send_json(200, {"success": True, "count": len(tickets), "tickets": tickets})

            elif path == "/api/import-excel-membres":
                ws = wb["Membres"]
                rows = list(ws.values)[1:]
                membres = [{"id": str(r[0] or ""), "prenom": str(r[1] or ""), "nom": str(r[2] or ""),
                            "adresse": str(r[3] or ""), "localite": str(r[4] or ""),
                            "email": str(r[5] or ""), "telephone": str(r[6] or ""),
                            "statut": str(r[7] or "Actif"), "dateAdhesion": str(r[8] or "")}
                           for r in rows if r[1] or r[2]]
                self.send_json(200, {"success": True, "count": len(membres), "membres": membres})

        except Exception as e:
            self.send_json(500, {"error": str(e)})


# ─── Démarrage ────────────────────────────────────────────────────────────────
def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except:
        return "127.0.0.1"

if __name__ == "__main__":
    PORT = 8000
    ip   = get_local_ip()

    print("╔══════════════════════════════════════════════════════╗")
    print("║           🍺  SERVEUR BAR LOCAL  v4.1  🍺            ║")
    print("║         ✓ Stockage multi-fichiers JSON               ║")
    print("║         ✓ Export Excel auto après modif membres      ║")
    print("╠══════════════════════════════════════════════════════╣")
    print(f"║  📁 {SCRIPT_DIR[:50]:<50} ║")
    print("╠══════════════════════════════════════════════════════╣")
    print(f"║  💻 PC  (caisse) → http://localhost:{PORT:<4}            ║")
    print(f"║  📱 GSM          → http://{ip}:{PORT}        ║")
    print("╠══════════════════════════════════════════════════════╣")
    print(f"║  🐍 Environnement : {'venv ✅':<45} ║" if IN_VENV else
          f"║  🐍 Environnement : {'système ⚠️  (venv recommandé)':<45} ║")
    print(f"║  📊 openpyxl      : {'✅ disponible':<45} ║" if EXCEL_OK else
          f"║  📊 openpyxl      : {'❌ manquant → pip install openpyxl':<45} ║")
    print("╠══════════════════════════════════════════════════════╣")
    missing = [f for f in ["tables.html", "commande.html"] if not os.path.exists(f)]
    if missing:
        for f in missing:
            print(f"║  ❌ Manquant : {f:<40} ║")
    else:
        print("║  ✅ Fichiers HTML trouvés                             ║")
    print("╠══════════════════════════════════════════════════════╣")
    print("║  📊 Fichiers de données :                            ║")
    for fname, fpath in FILES.items():
        exists = "✅" if os.path.exists(fpath) else "🆕"
        print(f"║    {exists} {os.path.basename(fpath):<48} ║")
    try:
        d = load_all()
        print(f"║  → {len(d.get('tables',[]))} tables · "
              f"{len(d.get('articlesData',{}).get('articles',[]))} articles · "
              f"{len(d.get('membres',[]))} membres · "
              f"{len(d.get('rapports',[]))} rapports ║")
    except Exception as e:
        print(f"║  ⚠️  {str(e)[:50]:<50} ║")
    print("║                                                      ║")
    print("║  Ctrl+C pour arrêter                                 ║")
    print("╚══════════════════════════════════════════════════════╝")
    print()

    import subprocess, time
    if sys.platform == "win32":
        try:
            result = subprocess.run(f'netstat -aon | findstr :{PORT} | findstr LISTENING',
                                    shell=True, capture_output=True, text=True)
            for line in result.stdout.strip().splitlines():
                parts = line.split()
                if parts:
                    subprocess.run(['taskkill', '/F', '/PID', parts[-1]], capture_output=True)
        except Exception:
            pass
    else:
        subprocess.run(["fuser", "-k", f"{PORT}/tcp"], capture_output=True)
    time.sleep(0.5)

    HTTPServer.allow_reuse_address = True
    try:
        server = HTTPServer(("0.0.0.0", PORT), BarHandler)
    except OSError as e:
        print(f"\n  ❌ Port {PORT} toujours occupé : {e}")
        sys.exit(1)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n  ✅ Serveur arrêté. À bientôt !")
        sys.exit(0)
